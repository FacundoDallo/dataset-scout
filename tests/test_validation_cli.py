import pandas as pd
import pytest
from openpyxl import load_workbook

from dataset_scout.cli import load_env_file, main
from dataset_scout.harmonize.vocabulary import Vocabulary
from dataset_scout.validation import compare, sheet_path, wilson_interval

CFG = "config/microglia_aging.yaml"


def test_wilson_interval():
    low, high = wilson_interval(95, 100)
    assert 0.88 < low < 0.89 and 0.97 < high < 0.99
    assert wilson_interval(0, 0) == (None, None)


def test_compare_counts_the_four_outcomes():
    samples = pd.DataFrame(
        {"gsm": ["A", "B", "C", "D"], "age_group": ["old", None, "young", "old"],
         "age_months": [24.0, None, 3.0, None], "sex": ["male", None, "female", "male"],
         "tissue": ["hippocampus", None, None, "brain"], "cell_type": [None] * 4,
         "sample_type": ["tissue"] * 4, "age_raw": ["24 mo", None, "3 mo", "aged"],
         "sex_raw": None, "tissue_raw": None, "cell_type_raw": None, "source_name": None}
    )
    sheet = pd.DataFrame(
        {"gsm": ["A", "B", "C", "D", "Z"], "age_group": ["old", "young", "old", "unknown", "old"],
         "age_months": ["23", "", "", "", ""], "sex": ["male", "unknown", "female", "", ""],
         "tissue": ["Hippocampus", "", "", "", ""], "cell_type": ["", "", "", "", ""],
         "sample_type": ["", "", "", "", ""], "reviewer_notes": ["", "", "", "", ""]}
    )
    results, disagreements, not_found = compare(sheet, samples, Vocabulary.load())
    age = results.set_index("field").loc["age_group"]
    assert (age.n_correct, age.n_missed, age.n_wrong, age.n_invented) == (1, 1, 1, 1)
    assert results.set_index("field").loc["age_months", "n_correct"] == 1  # 23 vs 24 is within tolerance
    assert results.set_index("field").loc["tissue", "accuracy"] == 1.0  # capitalization does not matter
    assert not_found == 1
    assert set(disagreements["outcome"]) == {"missed", "wrong", "invented"}


def fill(path, value="old"):
    wb = load_workbook(path)
    ws = wb["Review"]
    column = [c.value for c in ws[1]].index("age_group") + 1
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=column, value=value)
    wb.save(path)


def test_review_cycle_through_the_cli(recorded, config, capsys):
    assert main(["review-sheet", CFG, "--n", "12"]) == 0
    path = sheet_path(config)
    sheet = pd.read_excel(path, sheet_name="Review", dtype=str)
    assert len(sheet) == 12 and sheet["gse"].nunique() == 4
    assert "age_months" in sheet.columns and "program" not in " ".join(sheet.columns)
    assert main(["validate", CFG]) == 2  # nothing filled in yet
    fill(path)
    assert main(["validate", CFG]) == 0
    assert (config.validation_dir / "results.json").exists()
    assert "Is the harmonizer right?" in (config.output_dir / "report.html").read_text(encoding="utf-8")
    assert main(["review-sheet", CFG]) == 2  # refuses to overwrite manual work
    assert main(["review-sheet", CFG, "--out", "validation/second.xlsx"]) == 0


def test_config_and_sql_commands(recorded, capsys):
    assert main(["config", CFG]) == 0
    assert "GSE[ETYP]" in capsys.readouterr().out
    assert main(["sql", CFG, "SELECT gse, score FROM v_usable_studies"]) == 0
    assert "GSE1000001" in capsys.readouterr().out


def test_run_command_in_replay_mode(recorded, capsys):
    assert main(["run", CFG, "--mode", "replay"]) == 0
    assert "2 warn, 0 fail" in capsys.readouterr().out


def test_missing_snapshot_gives_a_clear_exit_code(project, capsys):
    assert main(["run", CFG, "--mode", "replay"]) == 4
    assert "Replay stopped" in capsys.readouterr().err


def test_bad_config_gives_a_clear_exit_code(project, capsys):
    (project / "bad.yaml").write_text("name: x\n", encoding="utf-8")
    assert main(["config", "bad.yaml"]) == 2
    assert "query.terms" in capsys.readouterr().err


def test_env_file(project, monkeypatch):
    monkeypatch.delenv("NCBI_EMAIL", raising=False)
    (project / ".env").write_text('# comment\nNCBI_EMAIL="me@example.org"\n', encoding="utf-8")
    assert load_env_file() == ["NCBI_EMAIL"]
    import os

    assert os.environ["NCBI_EMAIL"] == "me@example.org"
    monkeypatch.delenv("NCBI_EMAIL")


@pytest.mark.parametrize("args", [["--version"]])
def test_version(args, capsys):
    with pytest.raises(SystemExit):
        main(args)
    assert "dataset-scout" in capsys.readouterr().out
