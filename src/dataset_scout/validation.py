"""Blind manual review: how often is the harmonizer right?

Automated checks (checks.py) tell whether a build is internally
consistent. They cannot tell whether "18 mo" really became 18 months or
whether "HIP" really meant hippocampus. Only a person reading the raw
metadata can, so this module supports a small, GLP-style verification:

1. `scout review-sheet` draws a random set of in-scope samples, spread
   across as many studies as possible, and writes a spreadsheet that shows
   the raw metadata only. The program's answers are hidden on purpose: a
   reviewer who sees them tends to agree with them (anchoring bias).
2. A person fills in the yellow columns with the correct values.
3. `scout validate` (and every later `scout run`) compares the two and
   reports agreement per field with a 95% confidence interval, plus a file
   listing every disagreement for error analysis.

Every reviewed value lands in one of four outcomes:

- correct:  same value, or both agree the metadata does not say
- wrong:    both give a value and they differ
- missed:   the reviewer found a value the program did not
- invented: the program gave a value the reviewer could not find
"""

from __future__ import annotations

import hashlib
import json
import logging
import math
import random
from collections.abc import Callable
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd

from . import __version__
from .config import ScoutConfig
from .db import DatabaseError, query, replace_table
from .harmonize.vocabulary import Vocabulary
from .text import clean_words

log = logging.getLogger(__name__)

SHEET_NAME = "Review"
UNKNOWN_WORDS = frozenset({"unknown", "not stated", "not reported", "none", "na", "n/a", "-", "?"})
AGE_GROUPS = ["developmental", "young", "middle", "old", "adult", "unknown"]
SEXES = ["male", "female", "mixed", "unknown"]
SAMPLE_TYPES = ["tissue", "isolated cells", "primary culture", "cell line", "iPSC-derived", "organoid", "not stated"]

CONTEXT_COLUMNS = [
    "gsm",
    "gse",
    "series_title",
    "sample_title",
    "source_name",
    "characteristics",
    "organism",
    "geo_link",
]
INPUT_COLUMNS = [
    "age_group",
    "age_months",
    "sex",
    "tissue",
    "cell_type",
    "sample_type",
    "reviewer_notes",
]
INPUT_HELP = {
    "age_group": "developmental / young / middle / old / adult, or unknown if the metadata does not say",
    "age_months": "Age in months as a number (3 weeks = 0.69). Leave 'unknown' if no number is given",
    "sex": "male / female / mixed, or unknown",
    "tissue": "Tissue or organ, e.g. hippocampus, brain, blood. 'unknown' if not stated",
    "cell_type": "Cell type, e.g. microglia, astrocyte. 'unknown' if not stated",
    "sample_type": "tissue / isolated cells / primary culture / cell line / iPSC-derived / organoid / not stated",
    "reviewer_notes": "Optional: anything ambiguous or surprising",
}
EXAMPLE_ROW = {
    "gsm": "GSM0000000 (example)",
    "gse": "GSE000000",
    "series_title": "Transcriptome of microglia from young and aged mice",
    "sample_title": "MG_old_rep2",
    "source_name": "FACS-sorted CD11b+ microglia, whole brain",
    "characteristics": "strain: C57BL/6J\nage: 22 mo\nSex: female",
    "organism": "Mus musculus",
    "geo_link": "",
    "age_group": "old",
    "age_months": "22",
    "sex": "female",
    "tissue": "brain",
    "cell_type": "microglia",
    "sample_type": "isolated cells",
    "reviewer_notes": "",
}


@dataclass(frozen=True)
class FieldSpec:
    name: str
    column: str
    kind: str  # "category", "term" or "number"
    none_value: str | None = None
    evidence: str | None = None


FIELDS: tuple[FieldSpec, ...] = (
    FieldSpec("age_group", "age_group", "category", evidence="age_raw"),
    FieldSpec("age_months", "age_months", "number", evidence="age_raw"),
    FieldSpec("sex", "sex", "category", evidence="sex_raw"),
    FieldSpec("tissue", "tissue", "term", evidence="tissue_raw"),
    FieldSpec("cell_type", "cell_type", "term", evidence="cell_type_raw"),
    FieldSpec("sample_type", "sample_type", "category", none_value="not stated", evidence="source_name"),
)


def sheet_path(config: ScoutConfig) -> Path:
    return config.validation_dir / "review_sheet.xlsx"


def results_path(config: ScoutConfig) -> Path:
    return config.validation_dir / "results.json"


def disagreements_path(config: ScoutConfig) -> Path:
    return config.validation_dir / "disagreements.csv"


# --------------------------------------------------------------- the sheet
def _characteristics_text(chars: pd.DataFrame) -> str:
    lines = []
    for _, row in chars.sort_values(["channel", "position"]).iterrows():
        prefix = "" if int(row["channel"]) == 1 else f"[channel {int(row['channel'])}] "
        key = row["key_raw"]
        value = row["value_raw"]
        lines.append(f"{prefix}{key}: {value}" if isinstance(key, str) and key else f"{prefix}{value}")
    return "\n".join(lines)


def draw_review_sample(database: Path, n: int, seed: int) -> pd.DataFrame:
    """Pick up to `n` in-scope expression samples, spread across studies (round robin)."""
    population = query(
        database,
        """
        SELECT ss.gse, st.title AS series_title, s.gsm, s.title AS sample_title, s.source_name, s.organism
        FROM samples AS s
        JOIN study_samples AS ss USING (gsm)
        JOIN studies AS st USING (gse)
        WHERE st.in_scope AND s.is_expression
        ORDER BY ss.gse, s.gsm
        """,
    )
    if population.empty:
        raise DatabaseError("There are no in-scope expression samples to review. Run `scout run` first.")
    population = population.drop_duplicates(subset="gsm", keep="first")
    rng = random.Random(seed)
    by_study: dict[str, list[dict[str, Any]]] = {}
    for record in population.to_dict("records"):
        by_study.setdefault(record["gse"], []).append(record)
    order = sorted(by_study)
    rng.shuffle(order)
    for gse in order:
        rng.shuffle(by_study[gse])
    picked: list[dict[str, Any]] = []
    while len(picked) < n and any(by_study[g] for g in order):
        for gse in order:
            if by_study[gse] and len(picked) < n:
                picked.append(by_study[gse].pop())
    rng.shuffle(picked)
    frame = pd.DataFrame(picked)

    chars = query(database, "SELECT gsm, channel, position, key_raw, value_raw FROM sample_characteristics")
    grouped = {gsm: g for gsm, g in chars.groupby("gsm")}
    frame["characteristics"] = [
        _characteristics_text(grouped[g]) if g in grouped else "" for g in frame["gsm"]
    ]
    frame["geo_link"] = "https://www.ncbi.nlm.nih.gov/geo/query/acc.cgi?acc=" + frame["gsm"]
    for column in INPUT_COLUMNS:
        frame[column] = ""
    return frame[CONTEXT_COLUMNS + INPUT_COLUMNS]


def sheet_has_input(path: Path) -> bool:
    if not path.exists():
        return False
    try:
        existing = read_review_sheet(path)
    except Exception:  # an unreadable file still deserves protection
        return True
    filled = existing[[c for c in INPUT_COLUMNS if c in existing.columns]].fillna("")
    return bool((filled.map(lambda v: str(v).strip() != "")).to_numpy().any())


def write_review_sheet(frame: pd.DataFrame, path: Path, config: ScoutConfig, seed: int) -> Path:
    """Write the blind review workbook (Instructions + Review sheets)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    path.parent.mkdir(parents=True, exist_ok=True)
    arial = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    title_font = Font(name="Arial", size=14, bold=True)
    input_fill = PatternFill("solid", start_color="FFF2B3")
    header_fill = PatternFill("solid", start_color="E4E2EE")
    thin = Side(style="thin", color="C9C6D6")
    wrap_top = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()
    info = wb.active
    info.title = "Instructions"
    info["A1"] = "Blind review of harmonized sample metadata"
    info["A1"].font = title_font
    lines = [
        f"Run configuration: {config.name}",
        f"Samples drawn: {len(frame)} (random seed {seed}, spread across studies)",
        "",
        "What to do",
        "1. Open the Review sheet. Each row is one GEO sample.",
        "2. Read the grey columns (raw metadata as submitted). Open geo_link if you need more context.",
        "3. Fill in the yellow columns with the correct values. Do not look at the program's output first.",
        "4. Write 'unknown' when the metadata does not say. Leave a cell empty only if you skipped it.",
        "5. Save the file and run: scout validate <config file>",
        "",
        "Column guide (yellow = you fill in)",
    ]
    for offset, text in enumerate(lines, start=3):
        cell = info.cell(row=offset, column=1, value=text)
        cell.font = bold if text in {"What to do", "Column guide (yellow = you fill in)"} else arial
    row = 3 + len(lines)
    for column in INPUT_COLUMNS:
        name_cell = info.cell(row=row, column=1, value=column)
        name_cell.font = bold
        name_cell.fill = input_fill
        info.cell(row=row, column=2, value=INPUT_HELP[column]).font = arial
        row += 1
    row += 1
    info.cell(row=row, column=1, value="Example of a completed row (not part of the review)").font = bold
    row += 1
    for col_index, column in enumerate(CONTEXT_COLUMNS + INPUT_COLUMNS, start=1):
        header = info.cell(row=row, column=col_index, value=column)
        header.font = bold
        header.fill = input_fill if column in INPUT_COLUMNS else header_fill
        value = info.cell(row=row + 1, column=col_index, value=EXAMPLE_ROW[column])
        value.font = arial
        value.alignment = wrap_top
    info.column_dimensions["A"].width = 34
    info.column_dimensions["B"].width = 90

    ws = wb.create_sheet(SHEET_NAME)
    columns = CONTEXT_COLUMNS + INPUT_COLUMNS
    widths = {
        "gsm": 13,
        "gse": 12,
        "series_title": 36,
        "sample_title": 26,
        "source_name": 30,
        "characteristics": 46,
        "organism": 14,
        "geo_link": 18,
        "age_group": 15,
        "age_months": 12,
        "sex": 11,
        "tissue": 18,
        "cell_type": 18,
        "sample_type": 17,
        "reviewer_notes": 30,
    }
    for col_index, column in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_index, value=column)
        cell.font = bold
        cell.fill = input_fill if column in INPUT_COLUMNS else header_fill
        cell.border = Border(bottom=thin)
        ws.column_dimensions[get_column_letter(col_index)].width = widths[column]
    for row_index, record in enumerate(frame.to_dict("records"), start=2):
        for col_index, column in enumerate(columns, start=1):
            value = record[column]
            cell = ws.cell(row=row_index, column=col_index, value=None if value == "" else value)
            cell.font = arial
            cell.alignment = wrap_top
            if column in INPUT_COLUMNS:
                cell.fill = input_fill
            if column == "geo_link" and value:
                cell.hyperlink = value
                cell.value = "open in GEO"
    ws.freeze_panes = "C2"
    last = len(frame) + 1

    def add_list(column: str, options: list[str]) -> None:
        letter = get_column_letter(columns.index(column) + 1)
        validation = DataValidation(
            type="list",
            formula1='"' + ",".join(options) + '"',
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Value not in the list",
            error="Pick a value from the list (use 'unknown' when the metadata does not say).",
        )
        ws.add_data_validation(validation)
        validation.add(f"{letter}2:{letter}{last}")

    add_list("age_group", AGE_GROUPS)
    add_list("sex", SEXES)
    add_list("sample_type", SAMPLE_TYPES)
    wb.save(path)
    return path


def read_review_sheet(path: Path) -> pd.DataFrame:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        frame = pd.read_excel(path, sheet_name=SHEET_NAME, dtype=str)
    else:
        frame = pd.read_csv(path, dtype=str, sep=None, engine="python", encoding="utf-8-sig")
    frame.columns = [str(c).strip() for c in frame.columns]
    missing = [c for c in ["gsm", *INPUT_COLUMNS[:-1]] if c not in frame.columns]
    if missing:
        raise ValueError(f"The review sheet is missing columns: {', '.join(missing)}")
    return frame


# ------------------------------------------------------------- the scoring
def wilson_interval(successes: int, total: int, z: float = 1.96) -> tuple[float | None, float | None]:
    """95% confidence interval for a proportion (Wilson score method)."""
    if total == 0:
        return None, None
    p = successes / total
    denominator = 1 + z * z / total
    centre = (p + z * z / (2 * total)) / denominator
    half = z * math.sqrt(p * (1 - p) / total + z * z / (4 * total * total)) / denominator
    return max(0.0, centre - half), min(1.0, centre + half)


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, float) and math.isnan(value)) or str(value).strip() == ""


def _normalizer(spec: FieldSpec, vocab: Vocabulary) -> Callable[[Any], Any]:
    matcher = vocab.tissues if spec.name == "tissue" else vocab.cell_types

    def normalize(value: Any) -> Any:
        if _blank(value):
            return None
        text = clean_words(value)
        if text in UNKNOWN_WORDS or text == clean_words(spec.none_value or ""):
            return None
        if spec.kind == "number":
            try:
                return float(text.replace(",", "."))
            except ValueError:
                return text
        if spec.kind == "term":
            term = matcher.find(text)
            return clean_words(term.canonical) if term else text
        return text

    return normalize


def _same(spec: FieldSpec, expected: Any, got: Any) -> bool:
    if spec.kind == "number" and isinstance(expected, float) and isinstance(got, float):
        return abs(expected - got) <= max(0.5, 0.10 * abs(expected))
    return expected == got


def compare(sheet: pd.DataFrame, samples: pd.DataFrame, vocab: Vocabulary) -> tuple[pd.DataFrame, pd.DataFrame, int]:
    """Return (per-field results, disagreements, rows not found in this run)."""
    merged = sheet.merge(samples, on="gsm", how="left", suffixes=("_expected", ""), indicator=True)
    not_found = int((merged["_merge"] == "left_only").sum())
    merged = merged[merged["_merge"] == "both"]
    results, disagreements = [], []
    for spec in FIELDS:
        normalize = _normalizer(spec, vocab)
        counts = {"correct": 0, "wrong": 0, "missed": 0, "invented": 0}
        expected_column = f"{spec.column}_expected"
        for _, row in merged.iterrows():
            if _blank(row[expected_column]):
                continue
            expected = normalize(row[expected_column])
            got = normalize(row[spec.column])
            if expected is None and got is None:
                outcome = "correct"
            elif expected is None:
                outcome = "invented"
            elif got is None:
                outcome = "missed"
            elif _same(spec, expected, got):
                outcome = "correct"
            else:
                outcome = "wrong"
            counts[outcome] += 1
            if outcome != "correct":
                evidence = row.get(spec.evidence) if spec.evidence else None
                disagreements.append(
                    {
                        "gsm": row["gsm"],
                        "field": spec.name,
                        "outcome": outcome,
                        "reviewer_value": row[expected_column],
                        "program_value": None if _blank(row[spec.column]) else row[spec.column],
                        "raw_evidence": None if _blank(evidence) else evidence,
                        "reviewer_notes": row.get("reviewer_notes"),
                    }
                )
        reviewed = sum(counts.values())
        low, high = wilson_interval(counts["correct"], reviewed)
        results.append(
            {
                "field": spec.name,
                "n_reviewed": reviewed,
                "n_correct": counts["correct"],
                "n_wrong": counts["wrong"],
                "n_missed": counts["missed"],
                "n_invented": counts["invented"],
                "accuracy": round(counts["correct"] / reviewed, 4) if reviewed else None,
                "ci_low": round(low, 4) if low is not None else None,
                "ci_high": round(high, 4) if high is not None else None,
            }
        )
    return (
        pd.DataFrame(results),
        pd.DataFrame(
            disagreements,
            columns=["gsm", "field", "outcome", "reviewer_value", "program_value", "raw_evidence", "reviewer_notes"],
        ),
        not_found,
    )


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def refresh_validation(config: ScoutConfig, database: Path) -> dict[str, Any] | None:
    """Score the review sheet against the current database, if a filled sheet exists."""
    path = sheet_path(config)
    if not sheet_has_input(path):
        return None
    sheet = read_review_sheet(path)
    samples = query(database, "SELECT * FROM samples")
    vocab = Vocabulary.load(config.vocabulary_path)
    results, disagreements, not_found = compare(sheet, samples, vocab)
    reviewed_rows = int(
        sheet[[c for c in INPUT_COLUMNS[:-1]]].fillna("").map(lambda v: str(v).strip() != "").any(axis=1).sum()
    )
    payload = {
        "config_name": config.name,
        "computed_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        "tool_version": __version__,
        "sheet": path.as_posix(),
        "sheet_sha256": _sha256(path),
        "rows_in_sheet": int(len(sheet)),
        "rows_reviewed": reviewed_rows,
        "rows_not_in_run": not_found,
        "fields": results.to_dict("records"),
    }
    config.validation_dir.mkdir(parents=True, exist_ok=True)
    results_path(config).write_text(json.dumps(payload, indent=2), encoding="utf-8")
    disagreements.to_csv(disagreements_path(config), index=False, encoding="utf-8")
    replace_table(database, "validation_results", results)
    if not_found:
        log.warning("  %s reviewed samples are not part of this run and were ignored", not_found)
    log.info("  validation: %s reviewed samples compared", reviewed_rows)
    return payload
